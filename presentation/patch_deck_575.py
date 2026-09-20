#!/usr/bin/env python3
"""patch_deck_575.py — migrate the talk deck from the 19 July 2026 566-book
taxonomy to the 20 September 2026 575-book growth run (run_20260920_k9_s5,
equivalence class 3273ea3e577fdc99, nlp_hash 92b9f2d2151f0ee7).

Successor to patch_deck_566.py. Same targeted, per-slide, paragraph-level
approach, with one important difference: renames are applied in **two passes
via sentinels**, because several new topic names collide with old ones —
T5's old name ("Social Systems and Second-Order Constructivism") is T2's new
name, T3's old name is T4's new name, and T4's old name ("Cybernetics of Self")
is a substring of T7's new name ("Cybernetics of Self and Reimagination of
Self"). A single-pass find/replace would double-apply and corrupt them.

Reads:  presentation/CyberneticsNLP_Talk_v4_566.pptx   (566 taxonomy)
Writes: presentation/CyberneticsNLP_Talk_v5_575.pptx   (non-destructive)

What it changes:
  - slide 4  : corpus headline 739·566 -> 755·575; chapters 7,349 -> 6,449;
               end year 2025 -> 2026; footnote "566 analysed" -> "575 analysed".
  - slide 5,24: narrative "566 books" -> "575 books".
  - slide 10 : API-cost line "566 books" -> "575 books".
  - slide 11 : stability 0.365 / 6-of-9 / T1-unstable -> 0.367 / 7-of-9 /
               none unstable, with T1 flagged as the uninterpreted residual.
               Bands are 09c thresholds (>=0.30 stable); log_pipeline_run.py
               counts 2 stable — KI-11, still an open inconsistency.
  - slide 13 : the 9 topic names AND their 7 shown top-word lines, positionally
               (T1..T9), plus the footer line naming T8/T9. The clusters
               RECOMBINED rather than permuted this time (two July topics
               merged, one split, one dispersed), so neither names nor keywords
               map across by position — see docs/decisions.md.
  - slide 14 : the four era -> dominant-topic lines, recomputed data-led from
               this run's era x dominant-topic distribution.

ERA CUT (decided 20 Sep): each era is computed over its own printed caption
range, so Foundational is <=1969 and Second-Order is 1960-1979. These overlap
across the 1960s — the captions themselves overlap ("1940s – 1960s" /
"1960s – 1970s"), and the alternative (a disjoint <=1959 Foundational) leaves
that era with only 10 books. The overlap does not change which topics lead:
T9/T6 lead on either cut.

  Foundational Era   (<=1969, n=41): T9 Cognition 20 · T6 Formal/Control 7
  Second-Order Wave  (1960-79, n=88): T9 Cognition 34 · T6 Formal/Control 16
  Social Scale       (1980-99, n=89): T9 Cognition 23 · T3 Management 20
  Diffusion          (2000+,  n=388): T2 Social Systems 85 · T8 Pol. Econ. 63

NOTE — the 19 July deck carried an editorial override on the Second-Order Wave
line (it led with Social Systems to match the era's Varela/autopoiesis identity
even though the raw counts led elsewhere). That override is NOT carried forward:
the successor topic (T2) holds only 3 of 57 books in the 1970s, and 20 Sep
direction was to go data-led. Consequence to be aware of when presenting: T9
leads three of the four eras, so the era slide now tells a flatter story than
the July version — and T9 is the topic with the weakest vocabulary signature
(largest topic, no July parent, no exclusive top words). See ROADMAP #32/#34.

Names T1..T9 (re-validated 20 Sep, provisional single-rater): Residual —
uninterpreted · Social Systems and Second-Order Constructivism · Management and
Organisational Cybernetics · Biological and Ecological Regulation: Homeostasis
& Allostasis · Cybernetics and Digital Culture · Formal Foundation and Control
Engineering · Cybernetics of Self and Reimagination of Self · Political Economy
of Cybernetics · Cognition and Cybernetics.

Usage:  python3 presentation/patch_deck_575.py   (run from repo root)
Requires: python-pptx.
"""
from pptx import Presentation

SRC = 'presentation/CyberneticsNLP_Talk_v4_566.pptx'
OUT = 'presentation/CyberneticsNLP_Talk_v5_575.pptx'


def replace_on_slide(pres, idx, old, new):
    """Replace old->new in any paragraph on slide idx. Paragraph-level:
    collapses the match to run 0 (preserving that run's formatting)."""
    hit = False
    for sh in pres.slides[idx].shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            full = "".join(r.text for r in para.runs)
            if old in full:
                newfull = full.replace(old, new)
                if para.runs:
                    para.runs[0].text = newfull
                    for r in para.runs[1:]:
                        r.text = ""
                hit = True
    return hit


def set_paragraph(pres, idx, marker, newtext):
    """Rewrite the whole paragraph containing `marker` to `newtext`."""
    for sh in pres.slides[idx].shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            full = "".join(r.text for r in para.runs)
            if marker in full:
                if para.runs:
                    para.runs[0].text = newtext
                    for r in para.runs[1:]:
                        r.text = ""
                return True
    return False


# ── Simple, non-colliding replacements ───────────────────────────────────────
# (slide_idx, old, new)
REPL = [
    (4,  "739 · 566", "755 · 575"),
    (4,  "7,349", "6,449"),
    (4,  "1954–2025", "1954–2026"),
    # slide 14 — title and Diffusion-era caption carry the same end year.
    # The corpus now runs to 2026 (four 2026 imprints, incl. The Shan-Shui City
    # and Market Cybernetics). The 1954 start is left alone deliberately: the
    # earliest book is actually 1942 (Bateson/Mead, Balinese Character), so
    # "1954" was already a framing choice rather than a data bound.
    (14, "How the Field Has Evolved: 1954–2025", "How the Field Has Evolved: 1954–2026"),
    (14, "2000s – 2025", "2000s – 2026"),
    (5,  "566 books", "575 books"),
    (10, "566 books", "575 books"),
    (24, "566 books, 70 years", "575 books, 70 years"),
    (11, "mean stability 0.365, 6/9 topics stable (T1 unstable, 0.145)",
         "mean stability 0.367, 7/9 topics stable, none unstable "
         "(T1 held back as an uninterpreted residual)"),
    (11, "5-seed stability complete (mean=0.365, 6/9 stable; T1 unstable)",
         "5-seed stability complete (mean=0.367, 7/9 stable; none unstable)"),
]

# ── Colliding replacements — applied in two passes via sentinels ─────────────
# Each entry: (slide_idx, old_string, sentinel, new_string)
SENTINEL = [
    # slide 13 — nine topic names, positional T1..T7 shown + T8/T9 in footer
    (13, "History of Information Age and Cybernetics", "§§N1§§",
         "Residual — uninterpreted"),
    (13, "Extensions and Exploration of Cybernetics", "§§N2§§",
         "Social Systems and Second-Order Constructivism"),
    (13, "Biological and Ecological Regulation: Homeostasis & Allostasis", "§§N3§§",
         "Management and Organisational Cybernetics"),
    (13, "Cybernetics of Self", "§§N4§§",
         "Biological and Ecological Regulation: Homeostasis & Allostasis"),
    (13, "Social Systems and Second-Order Constructivism", "§§N5§§",
         "Cybernetics and Digital Culture"),
    (13, "Foundations of Cybernetics", "§§N6§§",
         "Formal Foundation and Control Engineering"),
    (13, "Management and Organisational Cybernetics", "§§N7§§",
         "Cybernetics of Self and Reimagination of Self"),
    (13, "T8 Control and Feedback Systems", "§§N8§§",
         "T8 Political Economy of Cybernetics"),
    (13, "T9 Digital Arts, Architecture, Design and Posthumanism", "§§N9§§",
         "T9 Cognition and Cybernetics"),

    # slide 13 — per-topic top-word lines, positional T1..T7
    (13, "machine · computer · wiener · century · technology · american", "§§W1§§",
         "city · qian · chinese · water · xuesen · ancient"),
    (13, "voice · sound · music · qian · chinese · china", "§§W2§§",
         "social · communication · society · environment · meaning · object"),
    (13, "brain · cell · animal · organism · evolution · energy", "§§W3§§",
         "decision · management · organization · variety · environment · organisation"),
    (13, "person · feel · bateson · child · family · behavior", "§§W4§§",
         "cell · brain · animal · neuron · energy · organism"),
    (13, "social · communication · language · meaning · object · distinction", "§§W5§§",
         "computer · machine · technology · medium · cybernetic · body"),
    (13, "define · entropy · probability · theorem · equation · shall", "§§W6§§",
         "input · variable · equation · output · rate · define"),
    (13, "organization · management · social · decision · variety · cybernetic", "§§W7§§",
         "bateson · person · feel · family · child · tell"),

    # slide 14 — era -> dominant-topic lines (data-led; see ERA CUT above)
    (14, "Dominant: T8 Control and Feedback Systems · T7 Management and Organisational Cybernetics",
         "§§E1§§",
         "Dominant: T9 Cognition and Cybernetics · T6 Formal Foundation and Control Engineering"),
    (14, "Dominant: T5 Social Systems and Second-Order Constructivism · T8 Control and Feedback Systems",
         "§§E2§§",
         "Dominant: T9 Cognition and Cybernetics · T6 Formal Foundation and Control Engineering"),
    (14, "Dominant: T5 Social Systems and Second-Order Constructivism · T7 Management and Organisational Cybernetics",
         "§§E3§§",
         "Dominant: T9 Cognition and Cybernetics · T3 Management and Organisational Cybernetics"),
    (14, "Dominant: T5 Social Systems and Second-Order Constructivism · T9 Digital Arts, Architecture, Design and Posthumanism",
         "§§E4§§",
         "Dominant: T2 Social Systems and Second-Order Constructivism · T8 Political Economy of Cybernetics"),
]

# (slide_idx, marker_substring, new_full_paragraph_text)
PARA_SET = [
    (4, "GST dates from the 1950s",
     "* GST dates from the 1950s; corpus coverage of that period may be incomplete.  "
     "Findings are provisional — the collection is not exhaustive.  "
     "575 analysed (OCR-corrupt books excluded)."),
    (13, "9-topic solution",
     "9-topic solution — 7 shown above. Also: T8 Political Economy of Cybernetics · "
     "T9 Cognition and Cybernetics.  T1 is retained in the k=9 fit but not interpreted — "
     "nine topics, eight interpreted."),
]


def main():
    pres = Presentation(SRC)
    misses = []

    for idx, old, new in REPL:
        if replace_on_slide(pres, idx, old, new):
            print(f"  ok    slide {idx}: {old[:46]!r} -> {new[:38]!r}")
        else:
            misses.append((idx, old))
            print(f"  MISS  slide {idx}: {old[:64]!r}")

    # Pass 1: old -> sentinel.  Longest old-strings first, so that a short old
    # name that is a substring of a longer one (e.g. "Cybernetics of Self")
    # cannot consume the longer match before it is reached.
    print("\n  -- sentinel pass 1 (old -> sentinel) --")
    for idx, old, sent, _new in sorted(SENTINEL, key=lambda x: -len(x[1])):
        if replace_on_slide(pres, idx, old, sent):
            print(f"  ok    slide {idx}: {old[:52]!r} -> {sent}")
        else:
            misses.append((idx, old))
            print(f"  MISS  slide {idx}: {old[:64]!r}")

    # Pass 2: sentinel -> new
    print("\n  -- sentinel pass 2 (sentinel -> new) --")
    for idx, _old, sent, new in SENTINEL:
        if replace_on_slide(pres, idx, sent, new):
            print(f"  ok    slide {idx}: {sent} -> {new[:46]!r}")
        else:
            misses.append((idx, sent))
            print(f"  MISS  slide {idx}: sentinel {sent} not found")

    print()
    for idx, marker, newtext in PARA_SET:
        if set_paragraph(pres, idx, marker, newtext):
            print(f"  ok    slide {idx}: paragraph rewrite ({marker[:30]!r})")
        else:
            misses.append((idx, marker))
            print(f"  MISS  slide {idx}: paragraph marker {marker[:30]!r}")

    pres.save(OUT)
    total = len(REPL) + 2 * len(SENTINEL) + len(PARA_SET)
    print(f"\nSaved {OUT}  ({total - len(misses)}/{total} edits applied)")
    if misses:
        print(f"WARNING: {len(misses)} strings not found — review those slides.")


if __name__ == '__main__':
    main()
